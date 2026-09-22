"""
file_processor.py — Multi-modal attachment parser for the Ultron Protocol.

Extracts structured text and Vision payloads from uploaded files:
  • PDF Documents (.pdf)         → Extracted text with page numbers
  • Excel Spreadsheets (.xlsx)   → Sheet summaries and Markdown data tables
  • CSV Data (.csv)              → Markdown table and summary metrics
  • Plain Text (.txt, .md, .py)  → Raw contents
  • Images (.png, .jpg, .webp)   → Base64 data for Claude Vision API
"""

from __future__ import annotations

import base64
import io
import mimetypes
from dataclasses import dataclass, field
from typing import Any


@dataclass
class ProcessedFile:
    filename: str
    file_type: str  # 'image', 'document', 'spreadsheet', 'text'
    text_content: str = ""
    # For Claude Vision image blocks:
    is_image: bool = False
    image_media_type: str = ""
    image_base64: str = ""
    file_size: int = 0


class FileProcessor:
    """Processes uploaded files and prepares them for Claude / Mem0 analysis."""

    SUPPORTED_IMAGE_TYPES = {
        "image/png": "image/png",
        "image/jpeg": "image/jpeg",
        "image/jpg": "image/jpeg",
        "image/webp": "image/webp",
        "image/gif": "image/gif",
    }

    @classmethod
    def process_file(cls, filename: str, content_bytes: bytes) -> ProcessedFile:
        """
        Parse raw bytes into a ProcessedFile structure.
        """
        lower_name = filename.lower()
        size = len(content_bytes)
        mime_type, _ = mimetypes.guess_type(filename)
        mime_type = mime_type or "application/octet-stream"

        # ── 1. Images (Claude Vision) ─────────────────────────────────────────
        if mime_type in cls.SUPPORTED_IMAGE_TYPES or any(
            lower_name.endswith(ext) for ext in [".png", ".jpg", ".jpeg", ".webp", ".gif"]
        ):
            media_type = cls.SUPPORTED_IMAGE_TYPES.get(mime_type, "image/png")
            b64_data = base64.b64encode(content_bytes).decode("utf-8")
            return ProcessedFile(
                filename=filename,
                file_type="image",
                is_image=True,
                image_media_type=media_type,
                image_base64=b64_data,
                text_content=f"[Attached Image: {filename} ({size / 1024:.1f} KB)]",
                file_size=size,
            )

        # ── 2. PDF Documents ──────────────────────────────────────────────────
        if lower_name.endswith(".pdf"):
            return cls._process_pdf(filename, content_bytes, size)

        # ── 3. Excel Spreadsheets ─────────────────────────────────────────────
        if lower_name.endswith(".xlsx") or lower_name.endswith(".xls"):
            return cls._process_excel(filename, content_bytes, size)

        # ── 4. CSV Files ──────────────────────────────────────────────────────
        if lower_name.endswith(".csv"):
            return cls._process_csv(filename, content_bytes, size)

        # ── 5. Word Documents (.docx / .doc) ─────────────────────────────────
        if lower_name.endswith(".docx") or lower_name.endswith(".doc"):
            return cls._process_word(filename, content_bytes, size)

        # ── 6. PowerPoint (.pptx) ─────────────────────────────────────────────
        if lower_name.endswith(".pptx"):
            return cls._process_pptx(filename, content_bytes, size)

        # ── 7. JSON / XML — treat as structured plain text ────────────────────
        if lower_name.endswith(".json") or lower_name.endswith(".xml"):
            return cls._process_text(filename, content_bytes, size)

        # ── 8. Plain Text & Code Files (fallback) ─────────────────────────────
        return cls._process_text(filename, content_bytes, size)

    @classmethod
    def _process_pdf(cls, filename: str, content_bytes: bytes, size: int) -> ProcessedFile:
        try:
            import pypdf

            reader = pypdf.PdfReader(io.BytesIO(content_bytes))
            num_pages = len(reader.pages)
            extracted_pages = []

            for i, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    extracted_pages.append(f"--- Page {i + 1} of {num_pages} ---\n{page_text.strip()}")

            if extracted_pages:
                full_text = f"=== DOCUMENT: {filename} ({num_pages} pages) ===\n\n" + "\n\n".join(extracted_pages)
            else:
                full_text = f"=== DOCUMENT: {filename} ({num_pages} pages) ===\n[No extractable text found in PDF]"

            return ProcessedFile(
                filename=filename,
                file_type="document",
                text_content=full_text,
                file_size=size,
            )
        except Exception as e:
            return ProcessedFile(
                filename=filename,
                file_type="document",
                text_content=f"[Error reading PDF {filename}: {e}]",
                file_size=size,
            )

    @classmethod
    def _process_excel(cls, filename: str, content_bytes: bytes, size: int) -> ProcessedFile:
        try:
            import pandas as pd

            excel_file = io.BytesIO(content_bytes)
            xl = pd.ExcelFile(excel_file)
            sheet_names = xl.sheet_names

            sections = [f"=== SPREADSHEET: {filename} (Sheets: {', '.join(sheet_names)}) ==="]

            for sheet in sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet)
                rows, cols = df.shape
                sections.append(f"\n--- Sheet: {sheet} ({rows} rows, {cols} columns) ---")

                if rows == 0:
                    sections.append("[Empty Sheet]")
                    continue

                # Top rows preview as Markdown table (max 30 rows to preserve context)
                preview_df = df.head(30)
                md_table = preview_df.to_markdown(index=False)
                sections.append(md_table)

                if rows > 30:
                    sections.append(f"\n[... {rows - 30} more rows not shown ...]")

            return ProcessedFile(
                filename=filename,
                file_type="spreadsheet",
                text_content="\n".join(sections),
                file_size=size,
            )
        except Exception as e:
            return ProcessedFile(
                filename=filename,
                file_type="spreadsheet",
                text_content=f"[Error reading Excel file {filename}: {e}]",
                file_size=size,
            )

    @classmethod
    def _process_csv(cls, filename: str, content_bytes: bytes, size: int) -> ProcessedFile:
        try:
            import pandas as pd

            df = pd.read_csv(io.BytesIO(content_bytes))
            rows, cols = df.shape
            sections = [
                f"=== CSV DATA: {filename} ({rows} rows, {cols} columns) ===",
            ]

            preview_df = df.head(40)
            sections.append(preview_df.to_markdown(index=False))

            if rows > 40:
                sections.append(f"\n[... {rows - 40} more rows not shown ...]")

            return ProcessedFile(
                filename=filename,
                file_type="spreadsheet",
                text_content="\n".join(sections),
                file_size=size,
            )
        except Exception as e:
            return ProcessedFile(
                filename=filename,
                file_type="spreadsheet",
                text_content=f"[Error reading CSV file {filename}: {e}]",
                file_size=size,
            )

    @classmethod
    def _process_word(cls, filename: str, content_bytes: bytes, size: int) -> ProcessedFile:
        """Extract text from Word .docx files using python-docx."""
        try:
            import docx as _docx  # python-docx

            doc = _docx.Document(io.BytesIO(content_bytes))
            sections = [f"=== WORD DOCUMENT: {filename} ==="]

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                # Preserve heading structure
                if para.style.name.startswith("Heading"):
                    level = para.style.name.replace("Heading ", "").strip()
                    prefix = "#" * (int(level) if level.isdigit() else 2)
                    sections.append(f"\n{prefix} {text}")
                else:
                    sections.append(text)

            # Extract tables
            for i, table in enumerate(doc.tables):
                sections.append(f"\n--- Table {i + 1} ---")
                for row in table.rows:
                    cells = " | ".join(cell.text.strip() for cell in row.cells)
                    if cells.strip():
                        sections.append(f"| {cells} |")

            full_text = "\n".join(sections)
            if len(full_text) > 50000:
                full_text = full_text[:50000] + "\n\n[... Truncated for token limits ...]"

            return ProcessedFile(
                filename=filename,
                file_type="document",
                text_content=full_text,
                file_size=size,
            )
        except Exception as e:
            return ProcessedFile(
                filename=filename,
                file_type="document",
                text_content=f"[Error reading Word document {filename}: {e}]",
                file_size=size,
            )

    @classmethod
    def _process_pptx(cls, filename: str, content_bytes: bytes, size: int) -> ProcessedFile:
        """Extract slide text from PowerPoint .pptx files using python-pptx."""
        try:
            from pptx import Presentation  # python-pptx

            prs = Presentation(io.BytesIO(content_bytes))
            sections = [f"=== POWERPOINT: {filename} ({len(prs.slides)} slides) ==="]

            for i, slide in enumerate(prs.slides, 1):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())
                if slide_texts:
                    sections.append(f"\n--- Slide {i} ---")
                    sections.extend(slide_texts)

            full_text = "\n".join(sections)
            if len(full_text) > 50000:
                full_text = full_text[:50000] + "\n\n[... Truncated for token limits ...]"

            return ProcessedFile(
                filename=filename,
                file_type="document",
                text_content=full_text,
                file_size=size,
            )
        except ImportError:
            return ProcessedFile(
                filename=filename,
                file_type="document",
                text_content=f"[python-pptx not installed. Run: pip install python-pptx to enable PowerPoint support.]",
                file_size=size,
            )
        except Exception as e:
            return ProcessedFile(
                filename=filename,
                file_type="document",
                text_content=f"[Error reading PowerPoint {filename}: {e}]",
                file_size=size,
            )

    @classmethod
    def _process_text(cls, filename: str, content_bytes: bytes, size: int) -> ProcessedFile:
        try:
            text = content_bytes.decode("utf-8", errors="replace")
            # Cap extremely large plain text files to first ~50KB
            if len(text) > 50000:
                text = text[:50000] + "\n\n[... Remaining content truncated for token limits ...]"

            return ProcessedFile(
                filename=filename,
                file_type="text",
                text_content=f"=== FILE: {filename} ===\n\n{text}",
                file_size=size,
            )
        except Exception as e:
            return ProcessedFile(
                filename=filename,
                file_type="text",
                text_content=f"[Error reading file {filename}: {e}]",
                file_size=size,
            )
